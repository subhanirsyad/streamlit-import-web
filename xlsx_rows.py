"""Baca 1 tab xlsx jadi list of tuple -- nilainya bertipe sama kayak openpyxl iter_rows(values_only=True),
tapi hemat RAM: mode read_only (streaming) dibanding load penuh (file TAP_SV 56 MB: 4,1 GB -> <1 GB).

Dipakai import_video_bank_unified, import_live_streaming_unified, import_raw_data_sc_unified.
"""

import openpyxl


def read_rows(path, sheet_name=None, strict=False):
    """strict=True: error kalau sheet_name gak ada. Selain itu jatuh ke tab pertama."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            if strict:
                raise ValueError(f'Sheet "{sheet_name}" tidak ketemu. Sheet yang ada: {wb.sheetnames}')
            sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        # export TikTok sering punya tag <dimension> salah/kosong; read_only percaya tag itu dan bisa
        # kepotong jadi 1 baris -> baca semua baris apa adanya
        ws.reset_dimensions()
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]  # baris kosong bisa datang sebagai list
    finally:
        wb.close()
    # tanpa dimensi, baris bisa lebih pendek dari header -> samakan lebarnya (loader mengindeks langsung)
    width = max((len(r) for r in rows), default=0)
    return [r + (None,) * (width - len(r)) if len(r) < width else r for r in rows]
